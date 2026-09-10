"""用真实格式验证正文提取、字符覆盖和分段位置。"""

import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from docx import Document as WordDocument
from fastapi import HTTPException
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from knowledge import ProcessInput, convert_file, embed_segments, split_documents
from haystack import Document


class ProcessingTests(unittest.TestCase):
    """检查不同文档格式与中文分段边界。"""

    def test_character_coverage(self):
        """去除重叠后完整恢复原文，含中文、金额、代码及 emoji。"""
        for length, overlap in [(256, 0), (256, 30), (256, 200), (512, 200), (2048, 50)]:
            for text in [("合同金额 1,234.50 元，不含税。\n    code(x)\n\n🙂备注：不得退款！" * 120), "连续中文" * 3000, "\n\n".join(f"第 {i} 节：唯一标题\n" + (f"本节序号 {i}。正文保留完整。" * 30) for i in range(10))]:
                segments = split_documents([Document(content=text)], length, overlap)
                restored = ""
                for segment in segments:
                    content = segment["content"]
                    self.assertLessEqual(len(content), length)
                    self.assertEqual(segment["character_count"], len(content))
                    # 按原文偏移核对分段内容和字符覆盖范围。
                    start = max(0, len(restored) - overlap)
                    self.assertTrue(text.startswith(content, start), (length, overlap, start))
                    restored = text[:start] + content
                self.assertEqual(restored, text)

    def test_supported_formats(self):
        """验证各类文件的正文、工作表名称和 PDF 页码。"""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for suffix, text in {
                ".txt": "合同金额 1234.50 元。", ".markdown": "# 合同\n金额 1234.50 元。", ".htm": "<html><body><article><p>合同金额 1234.50 元，不得退款。</p></article></body></html>", ".md": "# 合同\n\n金额 **1234.50** 元。",
                ".html": "<html><body><article><h1>合同</h1><p>合同金额 1234.50 元，不得退款。</p></article></body></html>",
                ".json": json.dumps({"金额": "1234.50", "说明": "不得退款"}, ensure_ascii=False),
                ".csv": "项目,金额\n合同,1234.50\n",
            }.items():
                path = root / ("sample" + suffix)
                path.write_text(text)
                segments = split_documents(convert_file(path), 256, 50)
                self.assertTrue(segments, suffix)
                self.assertTrue(all(segment["page_number"] is None for segment in segments), suffix)
            word = WordDocument()
            word.add_paragraph("合同金额 1234.50 元。")
            word.save(root / "sample.docx")
            slides = Presentation()
            slide = slides.slides.add_slide(slides.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = "合同金额 1234.50 元。"
            slides.save(root / "sample.pptx")
            workbook = Workbook()
            workbook.active.title = "合同"
            workbook.active.append(["项目", "金额"])
            workbook.active.append(["0012", "1234.50"])
            workbook.save(root / "sample.xlsx")
            for suffix in [".docx", ".pptx", ".xlsx"]:
                documents = convert_file(root / ("sample" + suffix))
                self.assertIn("1234.50", "".join(doc.content or "" for doc in documents), suffix)
                segments = split_documents(documents, 256, 50)
                self.assertTrue(segments, suffix)
                if suffix == ".xlsx":
                    self.assertIn("0012", "".join(doc.content or "" for doc in documents))
                    self.assertEqual(segments[0]["source_label"], "合同")
            # 核验首尾及中间含空白页的 PDF 正文页码。
            from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
            pdf = PdfWriter()
            for text in ["", "Contract page two", "", "Amount 1234.50 page four", ""]:
                page = pdf.add_blank_page(612, 792)
                if not text:
                    continue
                font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
                page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): pdf._add_object(font)})})
                stream = DecodedStreamObject()
                stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
                page[NameObject("/Contents")] = pdf._add_object(stream)
            pdf.write(root / "sample.pdf")
            segments = split_documents(convert_file(root / "sample.pdf"), 256, 50)
            self.assertEqual([s["page_number"] for s in segments], [2, 4])

    def test_invalid_and_empty_files(self):
        """验证空白、损坏和加密原件的失败原因。"""
        with self.assertRaises(HTTPException) as empty:
            split_documents([Document(content=" \n\t")], 256, 50)
        self.assertEqual(empty.exception.detail["code"], "empty_content")
        self.assertEqual(empty.exception.detail["stage"], "splitting")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "encrypted.pdf"
            writer = PdfWriter()
            writer.add_blank_page(612, 792)
            writer.encrypt("secret")
            writer.write(path)
            with self.assertRaises(HTTPException) as encrypted:
                convert_file(path)
            self.assertEqual(encrypted.exception.detail["code"], "encrypted_file")
            path = Path(directory) / "broken.json"
            path.write_text("{broken")
            with self.assertRaises(ValueError):
                convert_file(path)


class EmbeddingTests(unittest.TestCase):
    """检查分段向量生成的维度校验和失败原因。"""

    def setUp(self):
        """构造带向量配置的处理任务和固定分段。"""
        self.input = ProcessInput(
            organizationId=uuid4(), knowledgeBaseId=uuid4(), documentId=uuid4(), processingId=uuid4(),
            chunkLength=512, chunkOverlap=50, embeddingModelIdentifier="embedding-test", embeddingDimension=1024,
            embedding={"baseUrl": "https://models.test/v1", "apiKey": "test-key"},
        )
        self.segments = [{"position": 1, "content": "合同金额 1234.50 元。"}]

    def embedder(self, run):
        """构造记录构造参数并替换真实调用的向量组件工厂。"""
        captured = {}

        def factory(**kwargs):
            captured.update(kwargs)
            return type("Embedder", (), {"run": staticmethod(run)})()

        return factory, captured

    def test_dimension_and_failure(self):
        """维度不符和调用失败分别返回对应原因码。"""
        def mismatched(documents):
            return {"documents": [replace(document, embedding=[0.1] * 768) for document in documents]}

        with patch("knowledge.OpenAIDocumentEmbedder", self.embedder(mismatched)[0]):
            with self.assertRaises(HTTPException) as mismatch:
                embed_segments(self.input, self.segments)
        self.assertEqual(mismatch.exception.detail["code"], "embedding_dimension_mismatch")
        self.assertEqual(mismatch.exception.detail["stage"], "embedding")

        def failing(documents):
            raise RuntimeError("model unavailable")

        with patch("knowledge.OpenAIDocumentEmbedder", self.embedder(failing)[0]):
            with self.assertRaises(HTTPException) as failure:
                embed_segments(self.input, self.segments)
        self.assertEqual(failure.exception.detail["code"], "embedding_failed")
        self.assertEqual(failure.exception.detail["stage"], "embedding")

    def test_embedding_attached(self):
        """成功生成的向量按分段顺序写回。"""
        def embedded(documents):
            return {"documents": [replace(document, embedding=[float(index)] * 1024) for index, document in enumerate(documents)]}

        factory, captured = self.embedder(embedded)
        with patch("knowledge.OpenAIDocumentEmbedder", factory):
            embed_segments(self.input, self.segments)
        self.assertEqual(len(self.segments[0]["embedding"]), 1024)
        self.assertEqual(captured["model"], "embedding-test")
        self.assertEqual(captured["dimensions"], 1024)
        self.assertEqual(captured["batch_size"], 20)
        self.assertTrue(captured["raise_on_failure"])
